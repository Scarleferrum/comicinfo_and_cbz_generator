"""Read ComicInfo rows through the already-required openpyxl dependency."""

from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from .fields_v2_5 import EXCEL_HEADERS
from .model_v2_5 import ComicMetadata


def read_comics(path: str | Path, sheet_name: str | None = None) -> list[ComicMetadata]:
    return [comic for _, comic in read_comics_with_rows(path, sheet_name)]


def read_comics_with_rows(
    path: str | Path,
    sheet_name: str | None = None,
) -> list[tuple[int, ComicMetadata]]:
    """Read comics together with their one-based Excel row numbers."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, ())
        positions = {
            str(name).strip(): index
            for index, name in enumerate(headers)
            if name is not None
        }
        if "Title" not in positions:
            raise ValueError("Excel 缺少必要欄位: Title")

        comics = []
        for row_number, row in enumerate(rows, start=2):
            values = {
                name: row[positions[name]]
                if name in positions and positions[name] < len(row)
                else None
                for name in EXCEL_HEADERS
            }
            comic = ComicMetadata(values)
            if comic.get("Title"):
                comics.append((row_number, comic))
        return comics
    finally:
        workbook.close()


def create_template(path: str | Path) -> None:
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(EXCEL_HEADERS)
    workbook.save(path)
    workbook.close()


def ensure_excel_columns(
    path: str | Path,
    columns: tuple[str, ...],
) -> tuple[str, ...]:
    """Append missing headers to the first worksheet without replacing records."""
    path = Path(path)
    workbook = load_workbook(path)
    try:
        worksheet = workbook.worksheets[0]
        header_values = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        headers = {
            str(value).strip() for value in header_values if value is not None
        }
        missing = tuple(name for name in columns if name not in headers)
        if not missing:
            return ()

        last_header_column = max(
            (
                index
                for index, value in enumerate(header_values, start=1)
                if value is not None
            ),
            default=0,
        )
        next_column = last_header_column + 1
        source_header = worksheet.cell(row=1, column=max(1, last_header_column))
        while any(
            worksheet.cell(row=row, column=next_column + offset).value is not None
            for row in range(2, worksheet.max_row + 1)
            for offset in range(len(missing))
        ):
            next_column += 1
        for offset, name in enumerate(missing):
            target_header = worksheet.cell(row=1, column=next_column + offset)
            if source_header.has_style:
                _copy_cell_style(source_header, target_header)
            target_header.value = name
        workbook.save(path)
        return missing
    finally:
        workbook.close()


def append_comic(path: str | Path, data: ComicMetadata | dict) -> int:
    """Write one comic to the first wholly blank row of the first worksheet."""
    path = Path(path)
    if not path.exists():
        create_template(path)

    workbook = load_workbook(path)
    try:
        worksheet = workbook.worksheets[0]
        header_values = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        positions = {
            str(value).strip(): index + 1
            for index, value in enumerate(header_values)
            if value is not None
        }
        if not positions:
            for column, name in enumerate(EXCEL_HEADERS, start=1):
                worksheet.cell(row=1, column=column, value=name)
            positions = {
                name: index
                for index, name in enumerate(EXCEL_HEADERS, start=1)
            }

        missing = [name for name in EXCEL_HEADERS if name not in positions]
        if missing:
            raise ValueError("第一張工作表缺少必要欄位: " + ", ".join(missing))

        comic = data if isinstance(data, ComicMetadata) else ComicMetadata(data)
        font_source = _find_font_source(worksheet, positions["Title"])
        target_row = worksheet.max_row + 1
        scan_width = max(worksheet.max_column, max(positions.values()))
        for row_number in range(2, worksheet.max_row + 1):
            if all(
                worksheet.cell(row=row_number, column=column).value is None
                for column in range(1, scan_width + 1)
            ):
                target_row = row_number
                break

        for name in EXCEL_HEADERS:
            column = positions[name]
            target_cell = worksheet.cell(row=target_row, column=column)
            if not target_cell.has_style:
                source_cell = _find_style_source(
                    worksheet,
                    column,
                    target_row,
                )
                if source_cell is not None:
                    _copy_cell_style(source_cell, target_cell)
            if font_source is not None:
                target_cell.font = copy(font_source.font)
            target_cell.value = comic.get(name, "")

        source_height = (
            worksheet.row_dimensions[target_row - 1].height
            if target_row > 2
            else None
        )
        if (
            worksheet.row_dimensions[target_row].height is None
            and source_height is not None
        ):
            worksheet.row_dimensions[target_row].height = source_height

        workbook.save(path)
        return target_row
    finally:
        workbook.close()


def _find_style_source(worksheet, column: int, target_row: int):
    """Find the closest styled cell in this column, preferring prior records."""
    for row_number in range(target_row - 1, 1, -1):
        cell = worksheet.cell(row=row_number, column=column)
        if cell.has_style:
            return cell

    for row_number in range(target_row + 1, worksheet.max_row + 1):
        cell = worksheet.cell(row=row_number, column=column)
        if cell.has_style:
            return cell

    return None


def _find_font_source(worksheet, title_column: int):
    """Use the workbook's existing Title input font for all metadata cells."""
    for row_number in range(2, worksheet.max_row + 1):
        title_cell = worksheet.cell(row=row_number, column=title_column)
        if title_cell.value is not None and title_cell.has_style:
            return title_cell

    header_cell = worksheet.cell(row=1, column=title_column)
    return header_cell if header_cell.has_style else None


def _copy_cell_style(source, target) -> None:
    """Copy Excel's visible and number formatting without copying cell values."""
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)
